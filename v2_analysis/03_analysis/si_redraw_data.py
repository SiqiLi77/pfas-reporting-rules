"""Values for the redrawn Supporting Information Figures S1-S13 and reconstruction-record Figures R1-R2 (V2.03), with
the source of every value.

S1-S6, R1 and R2 are the figures of Supporting Information V1.62; S7-S13 are Figures 1-7 of manuscript V1.64. Values are read from
the tables of those two documents wherever a table reports them ("table" sources, parsed from the .docx files here, never
retyped). Elements that no table reports are handled in one of three ways:
  release - the plotted values published with the authors' code release v1.0.0 (author_release_v1_0_0/, commit and
            checksums in SOURCE.json): S1 rows 4-10, S2, S3b, R1, S10c-d, and the intervals of S5a, S9b and S11a,c;
            each series is checked against the tabulated values it must reproduce;
  printed - integers printed in the original figure or stated in the original text (for example S3b support counts);
  derived - arithmetic on table values (for example S12b biases, S8c segment counts).
The release series are also read off the original figure bitmaps (si_digitize_lib.py) as a cross-check, and S3b is
recomputed from the UCMR 5 file.

Run from 03_analysis/:  python3 si_redraw_data.py
Output: si_redraw_data.json, si_redraw_data.log
"""
import hashlib
import itertools
import json
import os
import re
import sys

import numpy as np
import pandas as pd
from docx import Document

import si_digitize_lib as dg

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
ORIG = os.path.join(HERE, "..", "01_manuscript", "original")
SI162 = os.path.join(ORIG, "CENMIXQ_EST_SUPPORTING_INFORMATION_V1_62.docx")
MS164 = os.path.join(ORIG, "CENMIXQ_EST_MANUSCRIPT_V1_64.docx")
BMP = os.path.join(HERE, "..", "02_review", "round6_figure_review", "extracted")
LOG = []
CHECKS = []


def log(*a):
    s = " ".join(str(x) for x in a)
    LOG.append(s)
    print(s)


def check(label, ok, detail=""):
    CHECKS.append((label, bool(ok), detail))
    log(f"  [{'ok' if ok else 'FAIL'}] {label} {detail}")


def sha(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


# ----------------------------------------------------------------------------------------------- table parsing
SI = Document(SI162)
T = [[[c.text.strip() for c in r.cells] for r in t.rows] for t in SI.tables]
SI_TEXT = "\n".join(p.text for p in SI.paragraphs)
NUM = r"[+\-−]?\d[\d,]*\.?\d*"


def num(s):
    return float(s.replace("−", "-").replace(",", "").replace("%", "").replace("+", "").strip())


def est(s):
    """'6.601 (4.106, 9.489)' -> (6.601, 4.106, 9.489); '100.00 (reference)' -> (100.0, None, None)."""
    m = re.match(rf"^\s*({NUM})\s*\(\s*({NUM}),\s*({NUM})\s*\)", s)
    if m:
        return tuple(num(x) for x in m.groups())
    m = re.match(rf"^\s*({NUM})", s)
    return (num(m.group(1)), None, None)


def interval(s):
    a, b = re.split(r",\s*", s.strip())
    return num(a), num(b)


def table(k, header=None):
    """SI V1.62 table k (1-based, document order, as in 02_review/round8_si_redraw/v162_v164_tables.txt)."""
    t = T[k - 1]
    if header:
        assert t[0][0] == header, (k, t[0][0], header)
    return t


def row(k, label, header=None, start=0):
    t = table(k, header)
    hits = [r for r in t[1:] if " | ".join(r[:start + 1]).endswith(label) or r[start] == label]
    assert len(hits) == 1, (k, label, len(hits))
    return hits[0]


SI_TABLE = {5: "Table S2, panel A"}


def src(k, label):
    return f"table #{k} ({SI_TABLE[k]}), row '{label}'"


OUT = {"_sources": {"SI_V1_62": {"file": os.path.basename(SI162), "sha256": sha(SI162)},
                    "MS_V1_64": {"file": os.path.basename(MS164), "sha256": sha(MS164)}}}


def bitmap(rel):
    p = os.path.join(BMP, rel)
    OUT["_sources"][rel] = {"sha256": sha(p)}
    return dg.load(p)


REL = os.path.join(HERE, "author_release_v1_0_0")
REL_SRC = json.load(open(os.path.join(REL, "SOURCE.json")))
OUT["_sources"]["release"] = {k: REL_SRC[k] for k in ("repository", "tag", "commit", "tarball_sha256")}
DIGITIZER = {}


def release(rel):
    """Plot-data CSV of the authors' code release, verified against the checksum recorded in SOURCE.json."""
    k = f"results/figures/{rel}"
    p = os.path.join(REL, k)
    assert sha(p) == REL_SRC["files"][k], rel
    OUT["_sources"][f"release:{rel}"] = {"sha256": REL_SRC["files"][k]}
    return pd.read_csv(p)


def rsrc(rel):
    return f"release: pfas-reporting-rules v1.0.0, results/figures/{rel}"


def cross_check(fig, unit, dig, rel):
    """Largest difference between values read off the original bitmap and the release values (informational)."""
    d = float(np.nanmax(np.abs(np.asarray(dig, float) - np.asarray(rel, float))))
    DIGITIZER[fig] = {"max_abs_difference": round(d, 4), "unit": unit}
    log(f"  digitizer vs release, {fig}: max |difference| {d:.4f} {unit}")


# =============================================================================================== S7 (V1.64 Fig. 1)
def s9():
    log("Figure S7")
    r3 = row(3, "UCMR 3 | Identifier matched", start=1)
    r5 = row(3, "UCMR 5 | Identifier matched", start=1)
    se1 = row(11, "SE1")
    n_u3, n_u5, n_u5old = (tuple(int(num(x)) for x in re.match(r"([\d,]+)/([\d,]+)", c).groups()) for c in se1[1:4])
    a = [{"rule": "UCMR 3 samples, UCMR 3 MRLs", "pct": num(r3[5]), "k": n_u3[0], "n": n_u3[1]},
         {"rule": "UCMR 5 samples, UCMR 5 MRLs", "pct": num(r5[5]), "k": n_u5[0], "n": n_u5[1]},
         {"rule": "UCMR 5 samples, UCMR 3 cutoffs", "pct": est(se1[3].split("(")[1])[0], "k": n_u5old[0], "n": n_u5old[1]}]
    for x in a:
        check(f"S7a {x['rule']} rate equals k/n", abs(100 * x["k"] / x["n"] - x["pct"]) < 0.001, f"{x['k']}/{x['n']} = {x['pct']}%")
    sys_ = []
    for lab, key in (("UCMR 3 native limits", "UCMR 3"), ("UCMR 5 native limits", "UCMR 5"), ("UCMR 5 at UCMR 3 cutoffs", "UCMR 5 old")):
        v = row(8, f"Matched PWSs with any K6≥2 event, {lab}")[1]
        k, n = (int(num(x)) for x in v.split(" of "))
        sys_.append({"rule": key, "systems": k, "of": n})
    m = re.search(r"ratios were ([\d.]+)% \((\d+)/(\d+)\) for UCMR 3 at native MRLs, ([\d.]+)% \(([\d,]+)/([\d,]+)\) for UCMR 5 at "
                  r"native MRLs, and ([\d.]+)% \((\d+)/(\d+)\) for UCMR 5 after", SI_TEXT)
    g = [num(x) for x in m.groups()]
    c = [{"pct": g[0], "k": int(g[1]), "n": int(g[2])}, {"pct": g[3], "k": int(g[4]), "n": int(g[5])},
         {"pct": g[6], "k": int(g[7]), "n": int(g[8])}]
    for x in c:
        check("S7c printed share equals k/n", abs(100 * x["k"] / x["n"] - x["pct"]) < 0.05, f"{x['k']}/{x['n']} = {x['pct']}%")
    pr = row(5, "Jurisdiction-equal, within-jurisdiction PWS-equal")
    d = {"native": est(pr[1]), "reclassification": est(pr[2]), "shared_cutoff": est(pr[3])}
    check("S7d native = reclassification + shared-cutoff (points)", abs(d["native"][0] - d["reclassification"][0] - d["shared_cutoff"][0]) < 0.0015,
          str(d))
    OUT["S7"] = {"a": a, "b": sys_, "c": c, "d": d,
                 "source": {"a": "table #3 (Table S1B) and #11 (Table S4B, SE1)", "b": "table #8 (Table S3B)",
                            "c": "printed: Text S2 ratios and counts", "d": src(5, "Jurisdiction-equal...")}}


# =============================================================================================== S1
S1_ROWS = ["Event-equal", "Pooled matched-PWS-equal", "Jurisdiction-equal, within-jurisdiction PWS-equal", "Metadata stable",
           "Normalized-name corroborated", "Same calendar month", "Same meteorological season", "Cyclic month distance <=1",
           "Cyclic month distance <=2", "Cyclic month distance <=3"]


def s1():
    log("Figure S1")
    im = bitmap("si_v162/image1.png")
    yt = dg.yticks(im, 840, 150, 1080)
    assert len(yt) == 10, yt
    panels = [("native", (841, 1260), (5, 10, 15), (184, 112, 144)), ("reclassification", (1357, 1776), (5, 10, 15), (64, 120, 144)),
              ("shared_cutoff", (1873, 2293), (-0.5, 0.0), (48, 136, 136))]
    rel = release("S1/robustness.csv")
    assert len(rel) == 10 and rel.analysis_or_restriction.iloc[3] == "Metadata stable"
    out = {k: [] for k, *_ in panels}
    worst, tab_err, dig_all, rel_all = 0.0, 0.0, [], []
    for col, ((key, (xl, xh), tv, rgb), pre) in enumerate(zip(panels, ("native", "reclassification", "common_cutoff"))):
        ticks = dg.xticks(im, 1087, xl - 2, xh + 2)
        ax = dg.Axis(ticks, tv)
        M = dg.near(im, rgb, 45)
        for i, y in enumerate(yt):
            mx, lo, hi = dg.hbar(M, y, xl, xh)
            dig = tuple(round(float(ax(v)), 4) for v in (mx, lo, hi))
            r = tuple(round(float(rel.loc[i, f"{pre}_{c}"]), 4) for c in ("estimate_pp", "ci95_low_pp", "ci95_high_pp"))
            dig_all.append(dig)
            rel_all.append(r)
            if i < 3:
                tab = est(row(5, S1_ROWS[i])[1 + col])
                worst = max(worst, max(abs(a - b) / ax.a for a, b in zip(dig, tab)))
                tab_err = max(tab_err, max(abs(a - b) for a, b in zip(r, tab)))
                out[key].append({"value": tab, "source": src(5, S1_ROWS[i]), "digitized_check": dig})
            else:
                out[key].append({"value": r, "source": rsrc("S1/robustness.csv"), "digitized_check": dig})
    check("S1 release reproduces the 9 tabulated estimates and 18 interval ends", tab_err <= 0.0005 + 1e-9,
          f"max |difference| {tab_err:.4f} pp (tables give 3 decimals)")
    check("S1 digitizer reproduces the 9 tabulated estimates and 18 interval ends", worst < 1.0, f"worst error {worst:.2f} px")
    cross_check("S1", "percentage points", dig_all, rel_all)
    supp = {r[0]: int(num(r[1])) for r in table(6)[1:]}
    n = [supp["Administrative identifier match"]] * 3 + [supp[k] for k in ("Metadata stable", "Normalized-name corroborated",
                                                                          "Same calendar month", "Same meteorological season",
                                                                          "Cyclic month distance <=1", "Cyclic month distance <=2",
                                                                          "Cyclic month distance <=3")]
    OUT["S1"] = {"rows": S1_ROWS, "pairs": n, **out, "digitizer_worst_px": round(worst, 2)}


# =============================================================================================== S2
def s2():
    log("Figure S2")
    im = bitmap("si_v162/image2.png")
    TEAL, PURP, BANDC = (55, 138, 140), (140, 121, 166), (225, 220, 232)
    alphas = np.round(np.arange(21) * 0.05, 2)
    tot = row(13, "Total")
    n_ev = int(num(tot[4]))
    end = {"a": 100 * int(num(tot[3])) / n_ev, "b": 100 * int(num(tot[1])) / n_ev}
    out = {"alpha": alphas.tolist()}
    yt_v = [100, 80, 60, 40, 20, 0]
    for name, xl, xh, spine in (("a", 226, 1110, 239), ("b", 1427, 2310, 1439)):
        ax = dg.Axis(dg.xticks(im, 781, xl, xh), [0, .5, 1])
        ay = dg.Axis(dg.yticks(im, spine, 100, 790), yt_v)
        T_, P, Bd = dg.near(im, TEAL, 22), dg.near(im, PURP, 16), dg.near(im, BANDC, 9)
        if name == "a":
            for M in (T_, P, Bd):
                M[150:215, 300:1100] = False
        x0, x2 = int(ax.px(0)) + 2, int(ax.px(1)) - 1
        X = np.arange(x0, x2)
        obs = np.array([160 + np.where(T_[160:779, c])[0].mean() if T_[160:779, c].any() else np.nan for c in X])
        ref = np.array([160 + np.where(P[160:779, c])[0].mean() if P[160:779, c].any() else np.nan for c in X])
        obs = np.interp(X, X[~np.isnan(obs)], obs[~np.isnan(obs)])
        ref = np.interp(X, X[~np.isnan(ref)], ref[~np.isnan(ref)])
        lo, hi = np.full(len(X), np.nan), np.full(len(X), np.nan)
        for j, c in enumerate(X):
            merged = []
            for a_, b_ in dg.runs(Bd[160:779, c]):
                if merged and a_ - merged[-1][1] <= 14:
                    merged[-1] = (merged[-1][0], b_)
                else:
                    merged.append((a_, b_))
            cand = [m for m in merged if m[0] - 3 <= ref[j] - 160 <= m[1] + 3]
            if cand:
                hi[j], lo[j] = 160 + cand[0][0], 160 + cand[0][1]

        def at(arr, a):
            d = np.abs(X - ax.px(a))
            k = d <= 2 if (d <= 2).any() else d <= d.min() + 2
            v = arr[k]
            v = v[~np.isnan(v)]
            return float(ay(v.mean())) if len(v) else np.nan

        O = [at(obs, a) for a in alphas]
        R = [at(ref, a) for a in alphas]
        L = [at(lo, a) for a in alphas]
        H = [at(hi, a) for a in alphas]
        start = 100.0 if name == "a" else 0.0
        O[0], R[0] = start, start
        O[-1] = end[name]
        f = {"a": "retained_multiple", "b": "collapsed_zero"}[name]
        r = release(f"S2/{f}.csv")
        r = r[(r.scheme == "jurisdiction_water") & (r.weighting == "event_equal")]
        assert np.allclose(r.alpha.values, alphas)
        keys = (("observed", "observed_percent"), ("reference_mean", "reference_mean_percent"),
                ("envelope_lo", "simultaneous_low_percent"), ("envelope_hi", "simultaneous_high_percent"))
        out[name] = {k: [None if np.isnan(v) else round(float(v), 4) for v in r[c].values] for k, c in keys}
        check(f"S2{name} release ends at the tabulated count at UCMR 3 cutoffs", abs(out[name]["observed"][-1] - end[name]) < 1e-3,
              f"{out[name]['observed'][-1]:.4f}% vs {end[name]:.4f}% (Table S5A)")
        cross_check(f"S2{name} observed and reference mean", "percentage points", O + R,
                    out[name]["observed"] + out[name]["reference_mean"])
        cross_check(f"S2{name} envelope", "percentage points", L[1:-1] + H[1:-1],
                    out[name]["envelope_lo"][1:-1] + out[name]["envelope_hi"][1:-1])
    integ = {k: float(np.trapezoid(np.array(out[p][s]) / 100, alphas)) for k, p, s in
             (("multiple_obs", "a", "observed"), ("multiple_ref", "a", "reference_mean"), ("none_obs", "b", "observed"),
              ("none_ref", "b", "reference_mean"))}
    tab = {"multiple_obs": num(row(18, "Event equal; multiple retained")[1]), "multiple_ref": num(row(18, "Event equal; multiple retained")[2]),
           "none_obs": num(row(18, "Event equal; none retained")[1]), "none_ref": num(row(18, "Event equal; none retained")[2])}
    for k in integ:
        check(f"S2 trapezoid integral of the release curve, {k}", abs(integ[k] - tab[k]) < 5e-6,
              f"{integ[k]:.6f} vs Table S7 {tab[k]:.6f}")
    out["integral_check"] = {k: [round(integ[k], 6), tab[k]] for k in integ}
    out["source"] = (rsrc("S2/retained_multiple.csv and collapsed_zero.csv") + " (jurisdiction-by-water-source reassignment, "
                     "event-equal); envelope = simultaneous 95% randomized-reference envelope, not defined at alpha 0 and 1")
    OUT["S2"] = out


# =============================================================================================== S3
S3B_PAIRS = ["PFHxA+PFPeA", "PFBS+PFPeA", "PFBS+PFHxA", "PFOA+PFPeA", "PFHxA+PFOA", "PFBS+PFOS", "PFOA+PFOS", "PFOS+PFPeA",
             "PFBA+PFPeA", "PFBS+PFOA", "PFHxA+PFOS", "PFBA+PFHxA"]
S3B_PRINTED_N = [1904, 1302, 1237, 1121, 1113, 1034, 1023, 994, 978, 972, 971, 873]
SIX = {"PFOA", "PFOS", "PFHxS", "PFNA", "PFHpA", "PFBS"}


def s3b_recompute():
    """First-sample pair support and re-reporting among UCMR 5 locations with an SE1 sample and a later complete sample."""
    from pilot_common import M533
    u = pd.read_parquet(os.path.join(HERE, "u5_pfas.parquet"))
    u = u[u.MethodID == "EPA 533"].copy()
    u["loc"] = u.PWSID + "|" + u.FacilityID + "|" + u.SamplePointID
    u["rep"] = (u.AnalyticalResultsSign == "=").astype(int)
    s = u.groupby(["loc", "SampleID"]).agg(nu=("Contaminant", "nunique"), date=("CollectionDate", "first"),
                                           se=("SampleEventCode", "first")).reset_index()
    s["date"] = pd.to_datetime(s.date)
    se1 = s[s.se == "SE1"]
    one = se1.groupby("loc").size()
    base = se1[se1["loc"].isin(one[one == 1].index) & (se1.nu == 25)].set_index("loc")
    W = u.pivot_table(index=["loc", "SampleID"], columns="Contaminant", values="rep", aggfunc="max").reindex(columns=M533).fillna(0)
    nxt = s[s["loc"].isin(base.index)].merge(base[["date"]].rename(columns={"date": "d0"}), left_on="loc", right_index=True)
    nxt = nxt[nxt.date > nxt.d0]
    first = nxt[nxt.date == nxt.groupby("loc").date.transform("min")]
    single = first.groupby("loc").filter(lambda g: len(g) == 1)
    good = single[single.nu == 25].set_index("loc")
    B = W.loc[list(zip(good.index, base.loc[good.index, "SampleID"]))].set_axis(good.index)
    A = W.loc[list(zip(good.index, good.SampleID))].set_axis(good.index)
    rows = {}
    for a, b in itertools.combinations(M533, 2):
        f = (B[a] == 1) & (B[b] == 1)
        if f.sum():
            rows[f"{a}+{b}"] = (int(f.sum()), int((f & (A[a] == 1) & (A[b] == 1)).sum()))
    return len(good), rows


def s3():
    log("Figure S3")
    a = []
    for panel, samples in (("Six", "First"), ("Six", "Two"), ("25", "First"), ("25", "Two")):
        r = [x for x in table(20)[1:] if x[0] == "Native" and x[1] == panel and x[2] == samples]
        assert len(r) == 1
        a.append({"panel": panel, "samples": samples, "label": est(r[0][3]), "pair": est(r[0][4])})
    im = bitmap("si_v162/image3.png")
    ax_a = dg.Axis(dg.xticks(im, 1277, 380, 1200), [0, 50, 100])
    ya = dg.yticks(im, 385, 250, 1270)
    TEAL, PINK = dg.near(im, (48, 136, 136), 30), dg.near(im, (184, 112, 144), 30)
    worst = 0.0
    for x, y in zip(a, ya):
        for key, M, dy in (("label", TEAL, -12), ("pair", PINK, 12)):
            if x[key][1] is None:
                continue
            r = dg.hbar(M, y + dy * 2.2, 390, 1190, half=10)
            worst = max(worst, abs(ax_a(r[0]) - x[key][0]) / ax_a.a)
    check("S3a digitizer reproduces the six tabulated point estimates", worst < 1.0, f"worst error {worst:.2f} px")
    ax_b = dg.Axis(dg.xticks(im, 1277, 1590, 2350), [0, 50, 100])
    yb = dg.yticks(im, 1596, 250, 1270)
    assert len(yb) == 12
    rel = release("S3/pair_recurrence.csv")
    assert [p.replace(" + ", "+") for p in rel.pair_or_group] == S3B_PAIRS
    b, dig = [], []
    for (_, r), p, y, n in zip(rel.iterrows(), S3B_PAIRS, yb, S3B_PRINTED_N):
        shared = set(p.split("+")) <= SIX
        mx, lo, hi = dg.hbar(TEAL if shared else PINK, y, 1600, 2340, half=14)
        dig.append([float(ax_b(v)) for v in (mx, lo, hi)])
        b.append({"pair": p, "shared_six": shared, "n_first_printed": n, "n_first": int(r.first_location_pairs),
                  "n_again": int(r.both_location_pairs),
                  "value": tuple(round(100 * float(r[c]), 3) for c in ("repeat_report_fraction", "ci95_lower", "ci95_upper"))})
    check("S3b release support counts equal the counts printed in the figure", [x["n_first"] for x in b] == S3B_PRINTED_N)
    cross_check("S3b", "percentage points", dig, [x["value"] for x in b])
    n_loc, rec = s3b_recompute()
    for x in b:
        nf, na = rec[x["pair"]]
        x["recomputed"] = {"n_first": nf, "n_again": na, "pct": round(100 * na / nf, 3)}
    top12 = sorted(rec, key=lambda k: -rec[k][0])[:12]
    check("S3b recomputed top-12 pairs are the printed pairs in the printed order", top12 == S3B_PAIRS)
    check("S3b recomputed paired cohort size", abs(n_loc - 25756) <= 2, f"{n_loc} vs 25,756 (Table S8A)")
    check("S3b recomputation from the UCMR 5 file reproduces the release counts and fractions",
          all((x["recomputed"]["n_first"], x["recomputed"]["n_again"]) == (x["n_first"], x["n_again"]) for x in b) and
          max(abs(x["recomputed"]["pct"] - x["value"][0]) for x in b) < 0.001,
          f"{sum(x['recomputed']['n_first'] == x['n_first'] for x in b)} of 12 support counts exact")
    OUT["S3"] = {"a": a, "b": b, "source": {"a": "table #20 (Table S8B), native rows",
                                           "b": rsrc("S3/pair_recurrence.csv") + "; support counts also printed in the "
                                                "figure; recomputed from u5_pfas.parquet as a check"}}


# =============================================================================================== S4
def s4():
    log("Figure S4")
    rows = []
    for lab in ("Location-proportional strata", "PWS equal within strata", "Jurisdiction equal", "Excluding Texas"):
        r = row(23, lab)
        rows.append({"weighting": lab, "six": est(r[1]), "25": est(r[2]), "change": est(r[3])})
        check(f"S4 change = 25 - six ({lab})", abs(rows[-1]["25"][0] - rows[-1]["six"][0] - rows[-1]["change"][0]) < 0.015)
    OUT["S4"] = {"rows": rows, "source": "table #23 (Table S9A, second part)"}


# =============================================================================================== S5, S9c
FLOORS = ["Native", "5", "10", "20", "40"]


def s5():
    log("Figure S5")
    rel = release("S5/independence_reference.csv")
    rows = []
    for f in FLOORS:
        r12 = row(27, f)
        g = rel[rel.floor_ng_L == (0 if f == "Native" else int(f))].set_index("component")
        ci = {k: tuple(round(float(g.loc[c, x]), 3) for x in ("estimate_pp", "ci95_lower", "ci95_upper"))
              for k, c in (("observed", "observed"), ("reference", "independent_reference"), ("difference", "observed_minus_independent"))}
        rows.append({"floor": f, "observed": ci["observed"], "reference": ci["reference"], "difference": est(r12[3]),
                     "observed_table_s11": est(row(26, f)[5])})
        check(f"S5 observed in Table S12 equals Table S11 ({f})", abs(num(r12[1]) - rows[-1]["observed_table_s11"][0]) < 1e-9)
        check(f"S5 difference = observed - reference ({f})",
              abs(rows[-1]["observed_table_s11"][0] - num(r12[2]) - rows[-1]["difference"][0]) < 0.0015)
        check(f"S5 release reproduces Table S12 ({f})",
              max(abs(ci["observed"][0] - num(r12[1])), abs(ci["reference"][0] - num(r12[2])),
                  *(abs(a - b) for a, b in zip(ci["difference"], rows[-1]["difference"]))) <= 0.0005 + 1e-9)
    OUT["S5"] = {"rows": rows, "source": rsrc("S5/independence_reference.csv") + " (observed and reference, with intervals from "
                                          "the same 1,999 draws as the differences); differences from table #27 (Table S12)"}


# =============================================================================================== S6
def s6():
    log("Figure S6")
    rows = []
    for lab, lab2 in (("Shared six", "Shared six"), ("Complete 25", "Complete 25"), ("Added trio", "Added trio"),
                      ("Other 16", "Other 16 added")):
        r = row(28, lab)
        lo_b, hi_b = (num(x) for x in r[4].split(" to "))
        e = row(29, lab2)
        rows.append({"panel": lab, "difference": est(r[3]), "bounds": (lo_b, hi_b), "lower_ci": est(e[1]), "upper_ci": est(e[2])})
        check(f"S6 bounds in Table S13 panels A and B agree ({lab})", abs(rows[-1]["lower_ci"][0] - lo_b) < 1e-9 and
              abs(rows[-1]["upper_ci"][0] - hi_b) < 1e-9)
    OUT["S6"] = {"rows": rows, "source": "tables #28-#29 (Table S13)"}


# =============================================================================================== R1
def s7():
    log("Figure R1")
    im = bitmap("si_v162/image7.png")
    ax = dg.Axis(dg.xticks(im, 855, 300, 2300), [0, .25, .5, .75, 1])
    ay = dg.Axis(dg.yticks(im, 311, 100, 850), [.2, .15, .1, .05, 0])
    check("R1 axis calibration residual", max(ax.resid, ay.resid) < 0.002, f"{ax.resid:.4f}, {ay.resid:.4f}")
    T_, P = dg.near(im, (52, 136, 140), 22), dg.near(im, (140, 120, 164), 18)
    T_[150:230, 500:1400] = False
    P[150:230, 500:1400] = False
    BD = (dg.near(im, (220, 224, 228), 7) | dg.near(im, (228, 240, 240), 7) | dg.near(im, (240, 236, 244), 7) |
          dg.near(im, (216, 224, 228), 5))
    LN = dg.near(im, (52, 136, 140), 60) | dg.near(im, (140, 120, 164), 60)
    AA = (im.min(axis=2) < 248) & (im.max(axis=2) > 130)
    alph = [.05, .10, .15, .20, .30, .35, .40, .45, .55, .60, .65, .70, .80, .85, .90, .95]
    xs = [ax.px(a) for a in alph]
    gau = []
    for x in xs:
        xi = int(round(x))
        win = P[200:850, xi - 11:xi + 12]
        w = win.sum(axis=1)
        r = np.where(w >= 7)[0]
        assert len(r), x
        gau.append(200 + (r * w[r]).sum() / w[r].sum())
    cols = np.arange(int(xs[0]) - 40, int(xs[-1]) + 40)
    cy = np.array([np.where(T_[200:850, c])[0].mean() + 200 if T_[200:850, c].any() else np.nan for c in cols])
    far = np.array([min(abs(c - x) for x in xs) > 16 for c in cols])
    segs = []
    for i in range(len(xs) - 1):
        k = (cols > xs[i]) & (cols < xs[i + 1]) & far & ~np.isnan(cy)
        segs.append(np.polyfit(cols[k], cy[k], 1))
    ss = [np.mean([np.polyval(s, x) for s in (segs[i - 1] if i else None, segs[i] if i < len(segs) else None) if s is not None])
          for i, x in enumerate(xs)]
    lo, hi = [], []
    for a, x, g in zip(alph, xs, gau):
        xi = int(round(x))
        cols_b = [xi + 3, xi + 4, xi + 5] if a == alph[0] else ([xi - 1, xi, xi + 1] if a <= 0.60 else [xi - 19, xi - 17, xi + 17, xi + 19])
        M = (BD | LN | AA) if a <= 0.60 else (BD | LN)
        res = []
        for c in cols_b:
            rr = dg.runs(M[200:850, c])
            merged = []
            for a_, b_ in rr:
                if merged and a_ - merged[-1][1] <= 3:
                    merged[-1] = (merged[-1][0], b_)
                else:
                    merged.append((a_, b_))
            cand = [m for m in merged if m[0] - 12 <= g - 200 <= m[1] + 12]
            if cand:
                res.append((200 + cand[0][0], 200 + cand[0][1]))
        hi.append(float(ay(np.mean([r[0] for r in res]))))
        lo.append(float(ay(np.mean([r[1] for r in res]))))
    rel = release("S7/threshold_path.csv")
    s_, g_ = (rel[rel.comparison == c].reset_index(drop=True) for c in ("factorized_minus_shared_shock", "factorized_minus_gaussian"))
    assert np.allclose(s_.alpha, alph) and np.allclose(g_.alpha, alph)
    out = {"alpha": alph, "shared_shock": [round(float(v), 5) for v in s_.estimate],
           "gaussian": [round(float(v), 5) for v in g_.estimate],
           "band_lo": [round(float(v), 5) for v in np.minimum(s_.simultaneous_ci95_low, g_.simultaneous_ci95_low)],
           "band_hi": [round(float(v), 5) for v in np.maximum(s_.simultaneous_ci95_high, g_.simultaneous_ci95_high)],
           "source": rsrc("S7/threshold_path.csv") + "; band = union of the two models' simultaneous 95% intervals"}
    cross_check("R1 model curves", "nats", [float(ay(v)) for v in ss] + [float(ay(v)) for v in gau],
                out["shared_shock"] + out["gaussian"])
    cross_check("R1 band below the axis top (0.20 nats)", "nats",
                [v for v, r in zip(lo + hi, out["band_lo"] + out["band_hi"]) if r < 0.2],
                [r for r in out["band_lo"] + out["band_hi"] if r < 0.2])
    check("R1 both series fall to about zero at the UCMR 3 end", max(abs(out["shared_shock"][-1]), abs(out["gaussian"][-1])) < 0.002)
    OUT["R1"] = out


# =============================================================================================== R2, S12a
S30_ROWS = ["Measured at native cutoffs", "Retained at older cutoffs", "Independent, fixed margins", "Shared shock, fixed margins",
            "Gaussian factor-2, fixed margins", "Count-prior diagnostic", "Training empirical composition",
            "Logistic, no auxiliaries", "Linear, no auxiliaries", "Nonlinear, no auxiliaries", "Logistic, three auxiliaries",
            "Linear, three auxiliaries", "Nonlinear, three auxiliaries"]
S30C_ROWS = ["Measured at native cutoffs", "Retained at older cutoffs", "Logistic, no auxiliaries", "Linear, no auxiliaries",
             "Nonlinear, no auxiliaries", "Logistic, three auxiliaries", "Linear, three auxiliaries", "Nonlinear, three auxiliaries"]


def s30(k, rows):
    return {lab: {"gw": est(row(k, lab)[1]), "sw": est(row(k, lab)[2]), "contrast": est(row(k, lab)[3])} for lab in rows}


def s8():
    log("Figure R2")
    nat, wa = s30(65, S30_ROWS), s30(66, S30C_ROWS)
    for lab, r in nat.items():
        check(f"R2/S12 national contrast = SW - GW ({lab})", abs(r["sw"][0] - r["gw"][0] - r["contrast"][0]) < 0.0151)
    sup = {f"{r[0]} {r[1]}": {"events": int(num(r[2])), "pws": int(num(r[3])), "multi": int(num(r[4]))} for r in table(64)[1:]}
    OUT["R2"] = {"national": nat, "washington": wa, "support": sup, "source": "tables #64-#66 (Table S15)"}


# =============================================================================================== S8 (V1.64 Fig. 2)
def s10():
    log("Figure S8")
    a = [{"native_count": r[0].replace("K6=", ""), "cells": [int(num(x)) for x in r[1:4]], "total": int(num(r[4]))}
         for r in table(13)[1:]]
    for r in a:
        check(f"S8a row sums ({r['native_count']})", sum(r["cells"]) == r["total"])
    b = [{"pattern": r[0], "events": int(num(r[1])), "cells": [int(num(x)) for x in r[2:5]]} for r in table(14)[1:]]
    for r in b:
        check(f"S8b row sums ({r['pattern'][:20]})", sum(r["cells"]) == r["events"])
    check("S8b rows sum to 829", sum(r["events"] for r in b) == 829)
    get = lambda lab: tuple(int(num(x)) for x in row(16, lab)[1].split("/"))  # noqa: E731
    z1, n1 = get("Six-panel-zero among 25-panel-positive events")
    single_multi, n_single = get("Multiple of 25 among six-panel-singleton events")
    zero_multi, _ = get("Multiple of 25 among six-panel-zero events")
    hid, n2 = get("Six-panel-zero or singleton among 25-panel-multiple events")
    multi6 = int(num(row(15, "Six PFAS at native limits")[1].split("/")[0]))
    c = [{"group": "≥1 of 25 PFAS", "n": n1, "segments": [z1, n_single, n1 - z1 - n_single]},
         {"group": "≥2 of 25 PFAS", "n": n2, "segments": [zero_multi, single_multi, n2 - zero_multi - single_multi]}]
    check("S8c >=2 shared segment equals the 829 six-panel multi-PFAS events in both groups",
          c[0]["segments"][2] == multi6 and c[1]["segments"][2] == multi6, str(c))
    check("S8c hidden bracket = 351/1,180", c[1]["segments"][0] + c[1]["segments"][1] == hid,
          f"{hid}/{n2} = {100 * hid / n2:.1f}%")
    OUT["S8"] = {"a": a, "b": b, "c": c, "c_hidden": [hid, n2],
                  "source": {"a": "table #13 (Table S5A)", "b": "table #14 (Table S5B)",
                             "c": "derived from tables #15-#16 (Table S6A-B): segments = counts given there"}}


# =============================================================================================== S9 (V1.64 Fig. 3)
S10_ANALYTES = ["PFBA", "PFPeA", "PFHxA", "HFPO-DA", "6:2 FTS", "4:2 FTS", "PFDoA", "PFPeS"]


def s11():
    log("Figure S9")
    a = {lab: {"gw": est(row(22, lab)[1]), "sw": est(row(22, lab)[2])} for lab in ("Six, native limits", "25, native limits")}
    b = [{"analyte": k, "gw": num(row(25, k)[1]), "sw": num(row(25, k)[2]), "difference": est(row(25, k)[3])} for k in S10_ANALYTES]
    total = sum(x["difference"][0] for x in b)
    eff = est(row(26, "Native")[5])[0]
    check("S9b contributions sum to the native panel effect", abs(total - eff) < 0.001, f"{total:.4f} vs {eff}")
    trio = sum(x["difference"][0] for x in b[:3])
    rel = release("Fig3/Fig3b_panel_attribution.csv").set_index("analyte_group")
    pp = {k: tuple(round(100 * float(rel.loc[k, c]), 4) for c in ("estimate", "ci95_lower", "ci95_upper")) for k in rel.index}
    check("S9b release reproduces the Table S10 contributions of PFBA, PFPeA and PFHxA",
          max(abs(a - c) for x in b[:3] for a, c in zip(pp[x["analyte"]], x["difference"])) <= 0.0001 + 1e-9)
    check("S9b release aggregate of the other 16 PFAS equals the sum of their Table S10 contributions",
          len(rel.loc["Other 16", "members"].split(";")) == 16 and abs(pp["Other 16"][0] - sum(x["difference"][0] for x in b[3:])) <= 0.0003,
          f"{pp['Other 16'][0]:.4f} vs {sum(x['difference'][0] for x in b[3:]):.4f}")
    other16 = {"difference": pp["Other 16"], "members": rel.loc["Other 16", "members"].split(";"),
               "source": rsrc("Fig3/Fig3b_panel_attribution.csv") + " (interval computed jointly, not by summing limits)"}
    c = [{"floor": f, "effect": est(row(26, f)[5])} for f in FLOORS]
    d = {}
    for lab in ("Groundwater", "Surface water"):
        r = row(21, lab)
        lab_ret, pair_ret = (num(x) for x in r[3].split(" / "))
        d[lab] = {"reference_positive": int(num(r[2])), "label": lab_ret, "pair": pair_ret}
    OUT["S9"] = {"a": a, "b": b, "b_trio": round(trio, 4), "b_other16": other16, "b_other_zero": 11, "c": c, "d": d,
                  "source": {"a": "table #22 (Table S9A)", "b": "table #25 (Table S10); the other 11 analytes contribute 0.0000",
                             "c": "table #26 (Table S11)", "d": "table #21 (Table S8C), six PFAS, two samples"}}


# =============================================================================================== S10 (V1.64 Fig. 4)
def s12():
    log("Figure S10")
    models = {"Independence": "Conditional independence", "Count prior": "Count-prior diagnostic", "Shared shock": "Shared-shock",
              "Gaussian factor-2": "Two-factor Gaussian"}
    a = {k: {"pattern": num(row(35, v)[1]), "count": num(row(35, v)[2]), "identity": num(row(35, v)[3])} for k, v in models.items()}
    for k, v in a.items():
        check(f"S10a pattern = count + identity ({k})", abs(v["pattern"] - v["count"] - v["identity"]) < 2e-6)
    b = {comp: {"Shared shock": est(row(36, lab)[1]), "Gaussian factor-2": est(row(36, lab)[2])}
         for comp, lab in (("Pattern", "Complete-pattern log-score reduction"), ("Count", "Count log-score reduction"),
                           ("Conditional identity", "Conditional-identity log-score reduction"))}
    im = bitmap("ms_v164/image4.png")
    axc = dg.Axis(dg.xticks(im, 1655, 250, 1040), [0.01, 0.1, 1, 10, 100], log=True)
    yc = dg.yticks(im, 263, 1000, 1650)
    axd = dg.Axis(dg.xticks(im, 1655, 1690, 2360), [0, 0.2, 0.4])
    check("S10 axis calibration residual", max(axc.resid, axd.resid) < 0.002, f"{axc.resid:.4f}, {axd.resid:.4f}")
    series = {"Observed": ((36, 52, 64), 20), "Independence": ((124, 132, 136), 22), "Shared shock": ((52, 136, 140), 22)}
    c = {}
    for k, (rgb, tol) in series.items():
        M = dg.near(im, rgb, tol)
        B = [dict(x, x=x["x"] + 268, y=x["y"] + 1020) for x in dg.blobs(M[1020:1650, 268:1030], 25)]
        v = [None] * 7
        for x in B:
            r = int(np.argmin([abs(x["y"] - y) for y in yc]))
            assert v[r] is None, (k, r)
            v[r] = float(axc(x["x"]))
        c[k] = v
    single = sum(num(r[2]) for r in table(42)[1:])
    zero = num(row(43, "All unreported")[2])
    relc = release("Fig4/Fig4c_count_distribution.csv")
    rc = {k: relc[relc.distribution == lab].sort_values("reported_pfas_count").probability_percent.round(5).tolist()
          for k, lab in (("Observed", "Observed"), ("Independence", "Conditional independence"), ("Shared shock", "Shared-shock"))}
    check("S10c release observed K=0 and K=1 equal Table R4C and the sum of Table R4B singletons",
          abs(rc["Observed"][0] - 100 * zero) < 5e-4 and abs(rc["Observed"][1] - 100 * single) < 5e-4,
          f"{rc['Observed'][0]:.4f}% and {rc['Observed'][1]:.4f}% vs {100 * zero:.4f}% and {100 * single:.4f}%")
    check("S10c release distributions sum to 100%", (relc.groupby("distribution").probability_percent.sum() - 100).abs().max() < 1e-9)
    cross_check("S10c (K=1-6)", "relative difference", [x / y - 1 for k in c for x, y in zip(c[k][1:], rc[k][1:]) if x is not None], [0])
    c_out = {"Observed": [100 * zero, 100 * single] + rc["Observed"][2:], "Independence": rc["Independence"],
             "Shared shock": rc["Shared shock"]}
    c_src = {"Observed": "K=0 and K=1 from tables #43 and #42 (Table R4C, R4B); K=2-6 " + rsrc("Fig4/Fig4c_count_distribution.csv"),
             "models": rsrc("Fig4/Fig4c_count_distribution.csv")}
    # panel d: patterns from the dot matrix, frequencies digitized
    filled = dg.near(im, (68, 124, 144), 26)[1020:1650, 1130:1680]
    ring = dg.near(im, (186, 198, 202), 16)[1020:1650, 1130:1680]
    names = ["PFBS", "PFHpA", "PFHxS", "PFNA", "PFOA", "PFOS"]
    obs_b = sorted([x for x in dg.blobs(dg.near(im, (36, 52, 64), 20)[1020:1650, 1690:2360], 25)], key=lambda x: x["y"])
    ind_b = sorted([x for x in dg.blobs(dg.near(im, (124, 132, 136), 22)[1020:1650, 1690:2360], 25)], key=lambda x: x["y"])
    ss_b = sorted([x for x in dg.blobs(dg.near(im, (52, 136, 140), 22)[1020:1650, 1690:2360], 25)], key=lambda x: x["y"])
    rows_y = [x["y"] + 21 for x in obs_b]
    dot_f = dg.blobs(filled, 30)
    col_centres = []
    for x in sorted(b["x"] for b in dot_f + dg.blobs(ring, 15)):
        if not col_centres or x - col_centres[-1][-1] > 30:
            col_centres.append([x])
        else:
            col_centres[-1].append(x)
    col_centres = [np.mean(g) for g in col_centres]
    patterns = []
    for y in rows_y:
        mem = [x for x in dot_f if abs(x["y"] - y) < 25]
        patterns.append(sorted({int(np.argmin([abs(x["x"] - cc) for cc in col_centres])) for x in mem}))
    d = []
    s18a = {r[0]: num(r[2]) for r in table(41)[1:]}
    check("S10d dot matrix has six constituent columns", len(col_centres) == 6, str([round(c) for c in col_centres]))
    for i, idx in enumerate(patterns):
        pat = " + ".join(names[j] for j in idx) if len(col_centres) == 6 else None
        d.append({"pattern": pat, "Observed": float(axd(obs_b[i]["x"] + 1690)), "Independence": float(axd(ind_b[i]["x"] + 1690)),
                  "Shared shock": float(axd(ss_b[i]["x"] + 1690))})
    check("S10d dot matrix identifies six two- or three-PFAS patterns", all(p and 2 <= p.count("+") + 1 <= 3 for p in (x["pattern"] for x in d)),
          str([x["pattern"] for x in d]))
    reld = release("Fig4/Fig4d_exact_multiple_patterns.csv")
    check("S10d release lists the dot-matrix patterns in the same order", list(reld.exact_pattern_label) == [x["pattern"] for x in d])
    cols = (("Observed", "observed_probability"), ("Independence", "conditional_independence_probability"),
            ("Shared shock", "shared_shock_probability"))
    cross_check("S10d", "percentage points", [[x[k] for k, _ in cols] for x in d],
                [[100 * r[c_] for _, c_ in cols] for _, r in reld.iterrows()])
    d = [{"pattern": r.exact_pattern_label, **{k: round(100 * float(r[c_]), 6) for k, c_ in cols}} for _, r in reld.iterrows()]
    errs = [abs(x["Observed"] - 100 * s18a[x["pattern"]]) for x in d if x["pattern"] in s18a]
    check("S10d release observed frequencies equal Table R4A", len(errs) >= 4 and max(errs) < 5e-5,
          f"{len(errs)} patterns tabulated; max |difference| {max(errs):.6f} percentage points")
    for x in d:
        if x["pattern"] in s18a:
            x["Observed_source"] = "table #41 (Table R4A)"
            x["Observed"] = 100 * s18a[x["pattern"]]
        else:
            x["Observed_source"] = rsrc("Fig4/Fig4d_exact_multiple_patterns.csv")
    OUT["S10"] = {"a": a, "b": b, "c": c_out, "c_source": c_src, "d": d,
                  "source": {"a": "table #35 (Table S14A)", "b": "table #36 (Table S14B)", "c": "see c_source",
                             "d": "patterns and model frequencies " + rsrc("Fig4/Fig4d_exact_multiple_patterns.csv") +
                                  "; observed frequencies from Table R4A where listed, otherwise from the release"}}


# =============================================================================================== S11 (V1.64 Fig. 5)
S25_INPUTS = ["No auxiliary", "Auxiliary count", "Individual flags", "Flags and concentrations"]
S26_COMP = ["None → count", "Count → flags", "Flags → flags and concentrations", "None → full three"]


def s13():
    log("Figure S11")
    t25 = table(55)
    a = {m: [num([r for r in t25[1:] if r[0] == m and r[1] == i][0][2]) for i in S25_INPUTS] for m in ("Linear", "Nonlinear")}
    t26 = table(56)

    def g26(m, comp, score):
        r = [x for x in t26[1:] if x[0] == m and x[1] == comp and x[2] == score]
        assert len(r) == 1, (m, comp, score)
        return {"value": num(r[0][3]), "pws": interval(r[0][4]), "jurisdiction": interval(r[0][5])}
    a_pairs = {m: {comp: g26(m, comp, "Pattern") for comp in S26_COMP} for m in ("Linear", "Nonlinear")}
    for m in a:
        check(f"S11a successive pattern differences equal Table R12 ({m})",
              all(abs((a[m][i] - a[m][i + 1]) - a_pairs[m][S26_COMP[i]]["value"]) < 2e-6 for i in range(3)))
    b = {m: {comp: g26(m, comp, "Identity, count 1–5") for comp in S26_COMP} for m in ("Linear", "Nonlinear")}
    t27 = table(57)

    def g27(m, score):
        r = [x for x in t27[1:] if x[0] == m and x[1] == score]
        assert len(r) == 1
        return {"none": num(r[0][2]), "three": num(r[0][3]), "value": num(r[0][4]), "pws": interval(r[0][5]), "county": interval(r[0][6])}
    c = {m: g27(m, "Pattern") for m in ("Logistic", "Linear", "Nonlinear")}
    d = {m: g27(m, "Identity, count 1–5") for m in ("Logistic", "Linear", "Nonlinear")}
    for m in c:
        check(f"S11c reduction = none - three ({m})", abs(c[m]["none"] - c[m]["three"] - c[m]["value"]) < 2e-6)
    ra = release("Fig5/Fig5a_national_pattern.csv")
    a_ci = {m: [tuple(float(v) for v in ra[(ra.model == m.lower()) & (ra.x == i)][["lower", "upper"]].iloc[0]) for i in range(4)]
            for m in a}
    err = max(abs(float(ra[(ra.model == m.lower()) & (ra.x == i)].y.iloc[0]) - a[m][i]) for m in a for i in range(4))
    check("S11a release absolute losses equal Table R11", err < 5e-7, f"max |difference| {err:.2e} nats")
    rc = release("Fig5/Fig5c_washington_pattern.csv").set_index("name")
    c_ci = {}
    for m, k in (("Logistic", "independent"), ("Linear", "linear"), ("Nonlinear", "nonlinear")):
        c_ci[m] = {lvl: (float(rc.loc[f"{pre}__{k}", "ci95_lower"]), float(rc.loc[f"{pre}__{k}", "ci95_upper"]))
                   for lvl, pre in (("none", "baseline"), ("three", "selected3"))}
        check(f"S11c release absolute losses equal Table R13 ({m})",
              max(abs(rc.loc[f"baseline__{k}", "estimate"] - c[m]["none"]), abs(rc.loc[f"selected3__{k}", "estimate"] - c[m]["three"])) < 5e-7)
    OUT["S11"] = {"a": a, "a_ci": a_ci, "a_paired": a_pairs, "b": b, "c": c, "c_ci": c_ci, "d": d, "inputs": S25_INPUTS,
                  "comparisons": S26_COMP,
                  "source": {"a": "table #55 (Table R11) and #56 (Table R12, pattern); intervals " + rsrc("Fig5/Fig5a_national_pattern.csv"),
                             "b": "table #56 (Table R12, identity count 1-5)",
                             "c": "table #57 (Table R13, pattern); intervals " + rsrc("Fig5/Fig5c_washington_pattern.csv"),
                             "d": "table #57 (Table R13, identity count 1-5)"}}


# =============================================================================================== S12 (V1.64 Fig. 6)
S14_ROWS = [("Measured at native cutoffs", "Measured"), ("Retained at older cutoffs", "Older-cutoff summary"),
            ("Independent, fixed margins", "Independence"), ("Shared shock, fixed margins", "Shared shock"),
            ("Gaussian factor-2, fixed margins", "Gaussian factor-2"), ("Count-prior diagnostic", "Count prior"),
            ("Training empirical composition", "Empirical composition"), ("Nonlinear, no auxiliaries", "Nonlinear, no auxiliaries")]


def s14():
    log("Figure S12")
    nat = OUT["R2"]["national"]
    a = [{"condition": lab, "key": k, **nat[k]} for k, lab in S14_ROWS]
    m = nat["Measured at native cutoffs"]
    b = [{"condition": lab, "number": i, "gw_bias": round(nat[k]["gw"][0] - m["gw"][0], 3), "sw_bias": round(nat[k]["sw"][0] - m["sw"][0], 3)}
         for i, (k, lab) in enumerate(S14_ROWS) if i]
    c = {}
    for k, src_ in (("gw", 73), ("sw", 74)):
        c[k] = [{"pfas": r[0], "observed": num(r[1]), "predicted": num(r[2]), "bias": est(r[3])} for r in table(src_)[1:]]
        for r in c[k]:
            check(f"S12c bias = predicted - observed ({k} {r['pfas']})", abs(r["predicted"] - r["observed"] - r["bias"][0]) < 0.0015)
    cnt = {r[0]: {"observed": num(r[1]), "predicted": num(r[2]), "bias": est(r[3])} for r in table(75)[1:]}
    OUT["S12"] = {"a": a, "b": b, "c": c, "expected_count": cnt,
                  "source": {"a": "table #65 (Table S15B)", "b": "derived: rate minus measured rate, table #65",
                             "c": "tables #73-#74 (Table R17A-B)"}}


# =============================================================================================== S13 (V1.64 Fig. 7)
def s15():
    log("Figure S13")
    ends = ["Functional-group state", "PFOA–PFOS", "PFHxS–PFOS"]

    def blk(k):
        return [{"candidate": r[0], "reference": r[1], **{e: est(r[2 + i]) for i, e in enumerate(ends)}} for r in table(k)[1:]]
    OUT["S13"] = {"endpoints": ends, "ucmr5": blk(80), "wa_pws": blk(81), "wa_county": blk(82),
                  "source": "tables #80-#82 (Table R19A-C)"}
    for x, y in zip(OUT["S13"]["wa_pws"], OUT["S13"]["wa_county"]):
        check(f"S13 Washington PWS and county blocks share point estimates ({x['candidate']})",
              all(x[e][0] == y[e][0] for e in ends))


def main():
    for f in (s9, s1, s2, s3, s4, s5, s6, s7, s8, s10, s11, s12, s13, s14, s15):
        f()
    n_ok = sum(ok for _, ok, _ in CHECKS)
    log(f"\n{n_ok} of {len(CHECKS)} checks passed")
    OUT["_checks"] = [{"check": a, "ok": b, "detail": c} for a, b, c in CHECKS]
    OUT["_digitizer_vs_release"] = DIGITIZER
    with open(os.path.join(HERE, "si_redraw_data.json"), "w") as f:
        json.dump(OUT, f, indent=1, ensure_ascii=False, default=float)
    with open(os.path.join(HERE, "si_redraw_data.log"), "w") as f:
        f.write("\n".join(LOG) + "\n")
    raise SystemExit(0 if n_ok == len(CHECKS) else 1)


if __name__ == "__main__":
    main()
